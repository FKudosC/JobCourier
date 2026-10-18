"""简历文件：保存原件并提取纯文本，支持 PDF / Word(docx) / 图片（OCR）。

按扩展名分发；图片走 RapidOCR（本地 onnxruntime 推理，首次调用加载模型），
解析组件缺失或文件损坏都抛 ``ValueError``，由配置页把原因展示给用户。
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from pypdf import PdfReader
from pypdf.errors import PdfReadError

from job.paths import USER_DATA_DIR
from job.utils import ocr

# 支持的图片扩展名（OCR）
IMAGE_SUFFIXES = frozenset({".jpg", ".jpeg", ".png", ".bmp", ".webp"})


class ResumePdf:
    """简历原件的保存与文本提取（类名保留，兼容既有调用）。"""

    # 简历存放目录
    DIR = USER_DATA_DIR / "resumes"

    # 生活照存放目录
    PHOTO_DIR = USER_DATA_DIR / "photos"

    @classmethod
    def save(cls, filename: str, data: bytes) -> Path:
        """把上传的简历存到 ``DIR``（同名覆盖），返回本地路径。"""
        cls.DIR.mkdir(parents=True, exist_ok=True)
        path = cls.DIR / Path(filename).name
        path.write_bytes(data)
        return path

    @classmethod
    def save_photo(cls, filename: str, data: bytes) -> Path:
        """把上传的生活照存到 ``PHOTO_DIR``（同名覆盖），返回本地路径。"""
        cls.PHOTO_DIR.mkdir(parents=True, exist_ok=True)
        path = cls.PHOTO_DIR / Path(filename).name
        path.write_bytes(data)
        return path

    @staticmethod
    def extract_text(filename: str, data: bytes) -> str:
        """按扩展名提取文本；不支持的格式或解析失败抛 ``ValueError``。"""
        suffix = Path(filename or "").suffix.lower()
        if suffix == ".pdf":
            return _from_pdf(data)
        if suffix == ".docx":
            return _from_docx(data)
        if suffix in IMAGE_SUFFIXES:
            return _from_image(data)
        raise ValueError("不支持的简历格式，请上传 PDF、Word（docx）或图片（jpg/png）")


def _clean(text: str) -> str:
    """压掉多余空行。"""
    return re.sub(r"\n\s*\n+", "\n", text).strip()


def _from_pdf(data: bytes) -> str:
    try:
        pages = PdfReader(io.BytesIO(data)).pages
        text = "\n".join(page.extract_text() or "" for page in pages)
    except PdfReadError as exc:
        raise ValueError("简历文件无法识别，请确认是有效的 PDF") from exc
    return _clean(text)


def _from_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - 依赖已随包安装
        raise ValueError("Word 解析组件未安装：请在项目目录运行 uv add python-docx") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Word 文件无法识别，请确认是 docx 格式（老 .doc 请先另存为 docx）") from exc
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    if not parts:
        return ""
    return _clean("\n".join(parts))


def _from_image(data: bytes) -> str:
    """图片 OCR：RapidOCR 本地推理，行按纵坐标排序后拼接。"""
    rows = ocr.lines(data)
    if not rows:
        return ""
    # 按行的纵坐标排序，同一行的再按横坐标从左到右
    ordered = sorted(rows, key=lambda row: (row[2], row[3]))
    return _clean("\n".join(text for text, _score, _top, _left in ordered))
