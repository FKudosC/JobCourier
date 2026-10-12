"""Tests for resume text extraction across formats (pdf / docx / image)."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from pypdf import PdfWriter

from job.utils import resume as resume_mod
from job.utils.resume import ResumePdf


def _blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("张三")
    document.add_paragraph("Python 后端开发，5 年经验")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "推荐系统重构"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extract_text_pdf():
    assert ResumePdf.extract_text("cv.pdf", _blank_pdf()) == ""


def test_extract_text_docx():
    text = ResumePdf.extract_text("简历.docx", _docx_bytes())
    assert "张三" in text
    assert "Python 后端开发" in text
    assert "推荐系统重构" in text  # 表格内容也要取到


def test_extract_text_corrupt_docx():
    with pytest.raises(ValueError, match="docx"):
        ResumePdf.extract_text("cv.docx", b"not a docx")


def test_extract_text_image_dispatches_to_ocr(monkeypatch: pytest.MonkeyPatch):
    called: list[bytes] = []

    def fake_ocr(data: bytes) -> str:
        called.append(data)
        return "张三\n13800000000"

    monkeypatch.setattr(resume_mod, "_from_image", fake_ocr)
    text = ResumePdf.extract_text("简历截图.png", b"fake-image-bytes")
    assert text == "张三\n13800000000"
    assert called == [b"fake-image-bytes"]


def test_extract_text_rejects_unsupported():
    with pytest.raises(ValueError, match="不支持的简历格式"):
        ResumePdf.extract_text("cv.doc", b"x")
    with pytest.raises(ValueError, match="不支持的简历格式"):
        ResumePdf.extract_text("cv", b"x")


def test_extract_text_rejects_corrupt_pdf():
    with pytest.raises(ValueError, match="简历文件无法识别"):
        ResumePdf.extract_text("cv.pdf", b"not a pdf")


def test_save_keeps_file_name_only(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(ResumePdf, "DIR", tmp_path)
    path = ResumePdf.save("../../evil/cv.pdf", b"%PDF")
    assert path == tmp_path / "cv.pdf"
    assert path.read_bytes() == b"%PDF"
